"""Таблица: CSV/XLSX из узла «Файл» или «Яндекс Диск» → строки, и обратно — строки → файл."""

from __future__ import annotations

import csv
import datetime
import io
import json
import os
from decimal import Decimal

import httpx

from helper import apply_variable_assignments, render_template
from schemas.workflow import SpreadsheetNode
from services.execution.actions.base import NodeAction
from services.execution.actions.data_transform import _extract_path
from services.execution.context import incoming_payload, parent_records
from services.media_files import file_output, load_binary, write_temp_bytes

MAX_ROWS = 5000
MAX_FILE_BYTES = 20 * 1024 * 1024
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_DELIMITERS = ",;\t|"


def _detect_format(requested: str, filename: str, mime: str, data: bytes) -> str:
    if requested != "auto":
        return requested
    ext = os.path.splitext(filename or "")[1].lower()
    if ext in {".xlsx", ".xlsm"} or "spreadsheetml" in mime:
        return "xlsx"
    if ext in {".csv", ".tsv", ".txt"} or mime.startswith("text/"):
        return "csv"
    return "xlsx" if data[:2] == b"PK" else "csv"


def _cell(value: object) -> object:
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def _trim(values: list[list[object]]) -> list[list[object]]:
    while values and all(item in (None, "") for item in values[-1]):
        values.pop()
    return values


def _read_xlsx(data: bytes, sheet: str) -> tuple[list[list[object]], str]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ValueError("Таблица: не удалось открыть файл XLSX") from None
    try:
        if sheet:
            if sheet not in workbook.sheetnames:
                raise ValueError(f"Таблица: в книге нет листа «{sheet}»")
            worksheet = workbook[sheet]
        else:
            worksheet = workbook.active
        values: list[list[object]] = []
        for row in worksheet.iter_rows(values_only=True):
            values.append([_cell(item) for item in row])
            if len(values) > MAX_ROWS:
                break
        return _trim(values), worksheet.title
    finally:
        workbook.close()


def _read_csv(data: bytes, delimiter: str) -> list[list[object]]:
    text = None
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Таблица: не удалось прочитать кодировку CSV")
    if not delimiter:
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=CSV_DELIMITERS).delimiter
        except csv.Error:
            delimiter = ","
    values: list[list[object]] = []
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        values.append(list(row))
        if len(values) > MAX_ROWS:
            break
    return _trim(values)


def _objects(values: list[list[object]]) -> tuple[list[str], list[dict]]:
    if not values:
        return [], []
    columns = [str(item).strip() if item not in (None, "") else f"col_{index + 1}" for index, item in enumerate(values[0])]
    rows = [
        {column: (row[index] if index < len(row) else None) for index, column in enumerate(columns)}
        for row in values[1:]
    ]
    return columns, rows


def _table(data: object) -> tuple[list[list[object]], bool]:
    """Строки для файла и признак, что первая строка — заголовки из ключей объектов."""
    if isinstance(data, dict):
        for key in ("rows", "values", "items", "data"):
            if isinstance(data.get(key), list):
                data = data[key]
                break
        else:
            data = [data]
    if not isinstance(data, list):
        raise ValueError("Таблица: для записи нужен массив строк или объектов")
    if not data:
        return [], False
    if all(isinstance(item, dict) for item in data):
        columns: list[str] = []
        for item in data:
            for key in item:
                if key not in columns:
                    columns.append(key)
        return [list(columns)] + [[item.get(column) for column in columns] for item in data], True
    return [list(item) if isinstance(item, (list, tuple)) else [item] for item in data], False


def _plain(value: object) -> object:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return json.dumps(value, ensure_ascii=False)


class SpreadsheetAction(NodeAction):
    def __init__(self, client: httpx.AsyncClient | None = None):
        self.client = client

    async def execute(self, node: SpreadsheetNode, context: dict) -> dict:
        if node.operation == "write":
            output = self._write(node, context)
        else:
            output = await self._read(node, context)
        apply_variable_assignments(node.assignments, output, context)
        return output

    async def _read(self, node: SpreadsheetNode, context: dict) -> dict:
        records = parent_records(node, context)
        incoming = file_output(records[0].get("output")) if records else None
        if incoming is None:
            raise ValueError("Таблица: подключите узел «Файл» или «Яндекс Диск» со скачанным файлом")
        if self.client is not None:
            data, mime, filename = await load_binary(incoming, self.client)
        else:
            async with httpx.AsyncClient(timeout=60) as client:
                data, mime, filename = await load_binary(incoming, client)
        if len(data) > MAX_FILE_BYTES:
            raise ValueError("Таблица: файл больше 20 МБ")
        file_format = _detect_format(node.file_format, filename, mime, data)
        sheet = ""
        if file_format == "xlsx":
            values, sheet = _read_xlsx(data, (node.sheet or "").strip())
        else:
            values = _read_csv(data, node.delimiter or "")
        truncated = len(values) > MAX_ROWS
        values = values[:MAX_ROWS]
        output: dict = {"filename": filename, "format": file_format, "truncated": truncated}
        if sheet:
            output["sheet"] = sheet
        if node.header_row:
            columns, rows = _objects(values)
            output.update({"columns": columns, "rows": rows, "row_count": len(rows)})
        else:
            output.update({"rows": values, "row_count": len(values)})
        return output

    def _write(self, node: SpreadsheetNode, context: dict) -> dict:
        payload = incoming_payload(node, context)
        variables = context.setdefault("variables", {})
        table, has_header = _table(_extract_path(payload, node.path))
        if not table:
            raise ValueError("Таблица: нет строк для записи")
        if len(table) > MAX_ROWS + 1:
            raise ValueError(f"Таблица: не больше {MAX_ROWS} строк")
        file_format = "csv" if node.file_format == "csv" else "xlsx"
        filename = os.path.basename(render_template(node.filename or "", payload, variables).strip()) or "table"
        if not filename.lower().endswith(f".{file_format}"):
            filename = f"{os.path.splitext(filename)[0] or 'table'}.{file_format}"
        if file_format == "xlsx":
            from openpyxl import Workbook

            workbook = Workbook(write_only=True)
            worksheet = workbook.create_sheet((node.sheet or "").strip()[:31] or "Sheet1")
            for row in table:
                worksheet.append([_plain(item) for item in row])
            buffer = io.BytesIO()
            workbook.save(buffer)
            data, mime = buffer.getvalue(), XLSX_MIME
        else:
            buffer = io.StringIO()
            writer = csv.writer(buffer, delimiter=(node.delimiter or ",")[:1])
            for row in table:
                writer.writerow(["" if item is None else _plain(item) for item in row])
            data, mime = buffer.getvalue().encode("utf-8-sig"), "text/csv"
        path = write_temp_bytes(data, f".{file_format}")
        return {
            "source": "file",
            "filename": filename,
            "mime": mime,
            "size": len(data),
            "path": path,
            "format": file_format,
            "row_count": len(table) - 1 if has_header or node.header_row else len(table),
        }
